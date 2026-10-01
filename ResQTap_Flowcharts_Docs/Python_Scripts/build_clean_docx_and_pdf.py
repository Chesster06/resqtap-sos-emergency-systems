import os
import glob
from PIL import Image
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
import win32com.client

BASE_DIR = r"c:\Users\Administrator\AndroidStudioProjects\ResQTap\ResQTap_Flowcharts_Docs"
PNG_DIR = os.path.join(BASE_DIR, "PNG_Diagrams")
DOCX_DIR = os.path.join(BASE_DIR, "Word_Documents_DOCX")
PDF_DIR = os.path.join(BASE_DIR, "PDF_Documents")

os.makedirs(DOCX_DIR, exist_ok=True)
os.makedirs(PDF_DIR, exist_ok=True)

FLOWCHARTS = [
    ("flowchart_1_overall_architecture_bw.png", "Flowchart_1_Overall_Architecture"),
    ("flowchart_2_auth_security_bw.png", "Flowchart_2_Authentication_Security"),
    ("flowchart_3_onetap_sos_bw.png", "Flowchart_3_OneTap_SOS"),
    ("flowchart_4_emergency_room_map_bw.png", "Flowchart_4_Emergency_Room_Hub"),
    ("flowchart_5_webrtc_call_bw.png", "Flowchart_5_WebRTC_Intercom"),
    ("flowchart_6_report_hospital_bw.png", "Flowchart_6_Reporting_Hospital"),
    ("flowchart_7_admin_portal_cloud_bw.png", "Flowchart_7_Admin_Portal"),
]

def setup_page(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)

def calculate_dimensions(img_path):
    img = Image.open(img_path)
    w_px, h_px = img.size
    
    # Available area on A4 landscape with 0.4 margin:
    # 11.69 - 0.8 = 10.89 in width
    # 8.27 - 0.8 = 7.47 in height
    # Leave small safety margin so Word never overflows onto page 2:
    w_avail = 10.60
    h_avail = 7.20
    
    scale = min(w_avail / float(w_px), h_avail / float(h_px))
    target_w = w_px * scale
    target_h = h_px * scale
    return target_w, target_h

def add_flowchart_image(doc, img_path):
    w_in, h_in = calculate_dimensions(img_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(0)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(w_in), height=Inches(h_in))

def build_individual_documents():
    print("--- Building 7 Individual DOCX Documents ---")
    generated_docx = []
    
    for filename, out_name in FLOWCHARTS:
        img_path = os.path.join(PNG_DIR, filename)
        if not os.path.exists(img_path):
            print(f"Warning: {filename} not found!")
            continue
            
        doc = docx.Document()
        setup_page(doc.sections[0])
        add_flowchart_image(doc, img_path)
        
        docx_path = os.path.join(DOCX_DIR, f"{out_name}.docx")
        doc.save(docx_path)
        print(f"Saved: {docx_path}")
        generated_docx.append((docx_path, os.path.join(PDF_DIR, f"{out_name}.pdf")))
        
    return generated_docx

def build_master_document():
    print("\n--- Building Master Combined DOCX Document (All 7 Flowcharts) ---")
    doc = docx.Document()
    
    for idx, (filename, out_name) in enumerate(FLOWCHARTS):
        img_path = os.path.join(PNG_DIR, filename)
        if not os.path.exists(img_path):
            continue
            
        if idx == 0:
            section = doc.sections[0]
            setup_page(section)
        else:
            section = doc.add_section()
            setup_page(section)
            
        add_flowchart_image(doc, img_path)
        print(f"  Added Page {idx + 1}: {out_name}")
        
    master_docx = os.path.join(DOCX_DIR, "ResQTap_All_Flowcharts.docx")
    doc.save(master_docx)
    print(f"Saved Master DOCX: {master_docx}")
    
    master_pdf = os.path.join(PDF_DIR, "ResQTap_All_Flowcharts.pdf")
    return master_docx, master_pdf

def export_all_to_pdf(doc_list):
    print("\n--- Exporting DOCX to PDF via Word COM ---")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    
    try:
        for docx_path, pdf_path in doc_list:
            doc = word.Documents.Open(docx_path)
            doc.ExportAsFixedFormat(pdf_path, 17) # 17 = wdExportFormatPDF
            doc.Close(False)
            print(f"Exported PDF: {pdf_path}")
    finally:
        word.Quit()

if __name__ == "__main__":
    # 1. Individual DOCX files
    indiv_list = build_individual_documents()
    
    # 2. Master DOCX file
    master_docx, master_pdf = build_master_document()
    
    # 3. Export all to PDF
    all_to_convert = indiv_list + [(master_docx, master_pdf)]
    export_all_to_pdf(all_to_convert)
    
    print("\nAll DOCX and PDF files generated successfully based on remaining diagram images!")
