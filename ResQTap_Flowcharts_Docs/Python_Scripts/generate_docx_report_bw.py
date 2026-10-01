import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

FONT_FAMILY = "Segoe UI"

# PURE MONOCHROME PALETTE
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_DARK_GRAY = RGBColor(40, 40, 40)
COLOR_MID_GRAY = RGBColor(100, 100, 100)
COLOR_WHITE = RGBColor(255, 255, 255)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders_bw(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="CCCCCC"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_bw(doc, text, title="NOTE / SPECIFICATION"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F5F5F5")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="000000"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"■  {title}\n")
    run_t.font.name = FONT_FAMILY
    run_t.font.size = Pt(10.5)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_BLACK
        
    run_b = p.add_run(text)
    run_b.font.name = FONT_FAMILY
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = COLOR_DARK_GRAY
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_1_bw(doc, text):
    h = doc.add_paragraph()
    format_paragraph(h, space_before=16, space_after=6)
    run = h.add_run(text)
    run.font.name = FONT_FAMILY
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return h

def add_heading_2_bw(doc, text):
    h = doc.add_paragraph()
    format_paragraph(h, space_before=12, space_after=4)
    run = h.add_run(text)
    run.font.name = FONT_FAMILY
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_BLACK
    return h

def add_heading_3_bw(doc, text):
    h = doc.add_paragraph()
    format_paragraph(h, space_before=8, space_after=2)
    run = h.add_run(text)
    run.font.name = FONT_FAMILY
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK_GRAY
    return h

def add_body_p_bw(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=0, space_after=6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_FAMILY
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
    r_body = p.add_run(text)
    r_body.font.name = FONT_FAMILY
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = COLOR_DARK_GRAY
    return p

def add_bullet_p_bw(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, space_before=0, space_after=3)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = FONT_FAMILY
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BLACK
    r_body = p.add_run(text)
    r_body.font.name = FONT_FAMILY
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_DARK_GRAY
    return p

def create_styled_table_bw(doc, headers, rows_data, col_widths=None):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_bw(tbl)
    
    # Header row: solid black with white bold text
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "000000")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        format_paragraph(p, 0, 0)
        for r in p.runs:
            r.font.name = FONT_FAMILY
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_WHITE
            
    # Data rows: white and light gray alternating
    for r_idx, row in enumerate(rows_data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_col = "F5F5F5" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            format_paragraph(p, 0, 0)
            for r in p.runs:
                r.font.name = FONT_FAMILY
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_BLACK
                if c_idx == 0:
                    r.font.bold = True
                    
    # Column widths
    if col_widths:
        for r in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)
                
    p_spacer = doc.add_paragraph()
    format_paragraph(p_spacer, 0, 4)
    return tbl

print("Monochrome Docx utilities defined successfully.")
